import copy
import base64
import json
from pathlib import Path
import tempfile
import unittest
from unittest.mock import patch

from harness import schemas
from harness.common import now, write_json, read_json, under
from harness.engine import Runner
from harness.quality import check_answer, check_slides, time_classification
from harness.materials import extract, select
from harness.exports import fixed_slide_count, text_pages
from harness.server import Manager
from harness.provider import Provider


PROBLEM = {'title': '逻辑练习', 'problem_statement': '1. 说明两个等量相加的结果。\n2. 解释结论成立的条件。',
           'questions': [{'id': '1', 'text': '说明两个等量相加的结果。'}, {'id': '2', 'text': '解释结论成立的条件。'}],
           'unclear': [], 'language': 'zh', 'word_min': 0, 'word_max': 0, 'slide_count': 0}


def answer(q):
    return {'question_id': q, 'answer': '结果是该量的两倍。', 'paragraphs': ['两个等量相加，可以提取相同的量，得到两倍关系。'],
            'tables': [], 'references': [], 'claims': [{'text': '两个等量相加得到两倍关系', 'kind': 'logic',
            'reference_ids': [], 'rationale': '由加法与乘法的定义得到。'}]}


class FakeProvider:
    calls = []
    fail = ''
    def __init__(self, config, directory, cancelled=lambda: False):
        self.usage = []
    def call(self, label, prompt, schema, images=(), search=False):
        self.calls.append(label)
        if self.fail == label:
            raise RuntimeError('模拟连接中断')
        self.usage.append({'label': label, 'seconds': .001, 'usage': {}})
        if label == 'transcription':
            return copy.deepcopy(PROBLEM)
        if label == 'overview':
            return {'title': '等量关系分析', 'summary': '两个等量相加得到两倍关系。', 'methods': '按题目顺序解释定义与适用条件。',
                    'lessons': ['将计算与适用条件分开讨论有助于清楚回答。']}
        if label.startswith('slides-'):
            return {'slides': [{'title': '等量关系', 'points': ['两个相同的量相加，结果为两倍。'],
                               'notes': '依据定义解释。', 'question_ids': ['1', '2'], 'reference_ids': []}]}
        if '-review-plan-' in label:
            return {'directions': [{'focus': '段落', 'reason': '原稿的适用条件可以更清楚', 'prompt': '明确本题等量条件并关联结论。'}]}
        for q in PROBLEM['questions']:
            if label.startswith(Runner.safe_id(q)):
                result = answer(q['id'])
                if 'revision' in label or 'fix' in label:
                    return {'answer': result, 'changes': [{'issue': '段落可以更直接', 'change': '保留结论并明确条件'}], 'blockers': []}
                return result
        raise AssertionError(label)


class PipelineTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.path = Path(self.temp.name)
        write_json(self.path / 'state.json', {'id': 'a'*32, 'started_at': now(), 'status': 'queued',
            'config': {}, 'metadata': {}, 'files': [], 'steps': {}, 'events': [], 'timings': [],
            'feedback': '', 'model_calls': [], 'answers': [], 'questions': []})
        FakeProvider.calls, FakeProvider.fail = [], ''
    def tearDown(self):
        self.temp.cleanup()
    def make_runner(self):
        return Runner(self.path, provider_factory=FakeProvider,
                      exporter=lambda p, c, r: {'report': 'test.docx', 'presentation': 'test.pptx'})
    def test_complete_two_rounds_per_question_and_original_statement(self):
        runner = self.make_runner()
        runner.execute()
        s = read_json(self.path/'state.json')
        self.assertEqual(s['status'], 'completed')
        self.assertEqual(s['problem_statement'], PROBLEM['problem_statement'])
        self.assertEqual((self.path/'Problem Statement.txt').read_text(), PROBLEM['problem_statement'])
        self.assertEqual(len([x for x in FakeProvider.calls if 'revision-1' in x]), 2)
        self.assertEqual(len([x for x in FakeProvider.calls if 'revision-2' in x]), 2)
        self.assertEqual([a['question_id'] for a in s['answers']], ['1', '2'])
        self.assertGreaterEqual(s['time']['elapsed_seconds'], 0)
    def test_resume_reuses_completed_question_work(self):
        failed = Runner.safe_id(PROBLEM['questions'][1])+'-draft'
        FakeProvider.fail=failed
        self.make_runner().execute()
        self.assertEqual(read_json(self.path/'state.json')['status'], 'error')
        first_draft = Runner.safe_id(PROBLEM['questions'][0])+'-draft'
        FakeProvider.fail=''
        self.make_runner().execute()
        self.assertEqual(FakeProvider.calls.count(first_draft), 1)
        self.assertEqual(read_json(self.path/'state.json')['status'], 'completed')
    def test_cancel_preserves_without_export(self):
        runner=self.make_runner()
        runner.stop.set()
        runner.execute()
        s=read_json(self.path/'state.json')
        self.assertEqual(s['status'],'paused')
        self.assertNotIn('export',s['steps'])
    def test_question_ids_do_not_collide(self):
        self.assertNotEqual(Runner.safe_id({'id':'Q1.1'}),Runner.safe_id({'id':'Q11'}))
    def test_repeated_question_removed_from_analysis_only(self):
        q=PROBLEM['questions'][0]
        a=answer(q['id']);a['paragraphs'].insert(0,q['text'])
        cleaned=Runner.clean_answer(q,a)
        self.assertNotIn(q['text'],cleaned['paragraphs'])
        self.assertEqual(q['text'],PROBLEM['questions'][0]['text'])


class QualityTests(unittest.TestCase):
    def test_unsupported_reference_fails_even_if_model_says_valid(self):
        a=answer('1')
        a['references']=[{'id':'r1','citation':'某作者. (2020). 某理论.', 'url':'', 'source_id':'S1',
                          'locator':'第1段','excerpt':'这是一段实际文件中根本不存在的文本'}]
        a['claims']=[{'text':'某命名理论成立','kind':'theory','reference_ids':['r1'],'rationale':''}]
        self.assertTrue(check_answer(a,{'id':'1'},[{'id':'S1','text':'另一段文本'}],{},False))
    def test_logical_claim_does_not_need_citation(self):
        self.assertEqual(check_answer(answer('1'),{'id':'1'},[],{},False),[])
    def test_body_style_is_checked_but_question_is_not_rewritten(self):
        a=answer('1');a['paragraphs']=['尚未形成证据链。']
        self.assertTrue(check_answer(a,{'id':'1','text':'讨论证据'},[],{},False))
    def test_slide_cannot_invent_numbers_or_omit_question(self):
        slides=[{'title':'利润999','points':['提升 999%'],'notes':'','question_ids':['2'],'reference_ids':[]}]
        self.assertTrue(check_slides(slides,[{'id':'1','text':'分析'}],[answer('1')]))
    def test_boundary_times_are_not_automatically_scored(self):
        for minutes in [180,600]:
            self.assertIn('补充口径',time_classification(minutes))
        self.assertIn('failure',time_classification(600.001))
        self.assertIn('90',time_classification(179.999))
    def test_schema_rejects_extras_and_wrong_boolean_integer(self):
        with self.assertRaises(ValueError):schemas.check({'value':True},schemas.obj(value=schemas.I))
        with self.assertRaises(ValueError):schemas.check({'value':'x','extra':1},schemas.obj(value=schemas.S))
    def test_path_escape_rejected(self):
        with self.assertRaises(ValueError):under('/tmp/run','../private')
    def test_docx_tables_preserved(self):
        from docx import Document
        with tempfile.TemporaryDirectory() as tmp:
            path=Path(tmp)/'source.docx'
            doc=Document();doc.add_paragraph('本次资料');t=doc.add_table(rows=2,cols=2)
            t.cell(0,0).text='项目';t.cell(0,1).text='金额';t.cell(1,0).text='收入';t.cell(1,1).text='30'
            doc.save(path)
            chunks,images=extract(path,'S1',Path(tmp)/'images')
            self.assertIn('收入 | 30',''.join(c['text'] for c in chunks))
    def test_gb18030_material_not_mistaken_for_utf16(self):
        with tempfile.TemporaryDirectory() as tmp:
            p=Path(tmp)/'data.csv';p.write_bytes('项目,收入\n甲,30'.encode('gb18030'))
            chunks,_=extract(p,'S1',Path(tmp)/'images')
            self.assertIn('项目,收入',''.join(c['text'] for c in chunks))
    def test_formula_without_cached_value_is_reported(self):
        import openpyxl
        with tempfile.TemporaryDirectory() as tmp:
            p=Path(tmp)/'formulas.xlsx';book=openpyxl.Workbook();book.active['A1']='=10+20';book.save(p);book.close()
            with self.assertRaisesRegex(ValueError,'公式没有保存计算结果'):
                extract(p,'S1',Path(tmp)/'images')
    def test_statement_pages_preserve_every_character_except_wrapping(self):
        statement='题目内容'+('ABC 中文内容 '*200)
        pages=text_pages(statement)
        self.assertEqual(''.join(''.join(p.splitlines()) for p in pages),statement)


class StorageAndProviderTests(unittest.TestCase):
    def test_api_key_is_not_saved_with_run(self):
        with tempfile.TemporaryDirectory() as tmp:
            manager=Manager(tmp)
            manager.start=lambda _:None
            payload={'files':[{'name':'../question.png','role':'question',
                              'data':base64.b64encode(b'test').decode()}],
                     'config':{'provider':'api','model':'example','api_key':'TEST_ONLY_PRIVATE_VALUE'},'metadata':{}}
            s=manager.create(payload)
            text=(Path(tmp)/s['id']/'state.json').read_text()
            self.assertNotIn('TEST_ONLY_PRIVATE_VALUE',text)
            self.assertEqual(manager.credentials[s['id']]['api_key'],'TEST_ONLY_PRIVATE_VALUE')
            self.assertTrue((Path(tmp)/s['id']/'uploads/0-question.png').exists())
            self.assertEqual(s['outputs'],{})
            manager.pool.shutdown()
    def test_unsupported_upload_is_rejected(self):
        with tempfile.TemporaryDirectory() as tmp:
            manager=Manager(tmp)
            with self.assertRaises(ValueError):
                manager.upload(Path(tmp),[{'name':'file.exe','role':'material','data':'AA=='}])
            manager.pool.shutdown()
    def test_finished_response_recovered_only_for_exact_prompt_and_schema(self):
        with tempfile.TemporaryDirectory() as tmp:
            path=Path(tmp)/'step';path.mkdir()
            schema=schemas.obj(value=schemas.S)
            (path/'prompt.txt').write_text('same input')
            write_json(path/'schema.json',schema)
            write_json(path/'last.json',{'value':'saved result'})
            provider=Provider({},tmp)
            with patch.object(provider,'_codex',side_effect=AssertionError('must not call model')):
                self.assertEqual(provider.call('step','same input',schema),{'value':'saved result'})
            with patch.object(provider,'_codex',return_value=({'value':'new result'},{})) as model:
                self.assertEqual(provider.call('step','changed input',schema),{'value':'new result'})
                self.assertEqual(model.call_count,1)


if __name__=='__main__':unittest.main()


class WordReportTests(unittest.TestCase):
    def test_repeated_part_stages_are_selected_without_other_part(self):
        from harness.word_format import select_part
        text='Assign-5-Part1 第一阶段\n要求甲\nAssign-5-Part2\n要求乙\nAssign-5-Part1 第二阶段\n要求丙'
        chosen=select_part(text, '1')
        self.assertIn('要求甲',chosen); self.assertIn('要求丙',chosen); self.assertNotIn('要求乙',chosen)
        with self.assertRaises(ValueError): select_part(text,'3')
    def test_prompt_markup_becomes_native_word_and_survives_save(self):
        from harness.exports import new_document
        from harness.word_format import markdown,heading,paragraph
        from docx import Document
        from docx.oxml.ns import qn
        with tempfile.TemporaryDirectory() as tmp:
            d=new_document();heading(d,'研究要求',1);heading(d,'任务',2)
            markdown(d,'# Role\n采用**真实资料**与*条件分析*，保留 [原论文](https://example.com/paper)。')
            p=paragraph(d.add_paragraph('中文 body'))
            file=Path(tmp)/'report.docx';d.save(file);saved=Document(file)
            self.assertEqual(saved.paragraphs[0].text,'一、研究要求')
            self.assertEqual(saved.paragraphs[1].text,'（一）任务')
            body=saved.paragraphs[-2]
            self.assertTrue(any(x.text=='真实资料' and x.bold for x in body.runs))
            self.assertTrue(any(x.text=='条件分析' and x.italic for x in body.runs))
            self.assertNotIn('**',body.text)
            self.assertEqual(saved.paragraphs[-1]._p.pPr.ind.get(qn('w:firstLineChars')),'200')
            self.assertEqual(saved.styles['Normal'].element.rPr.rFonts.get(qn('w:eastAsia')),'仿宋')
    def test_review_plan_precedes_changes_and_has_previous_context(self):
        runner=self.make_runner() if hasattr(self,'make_runner') else None
        # Flow ordering and substantive reasons are checked with the pipeline fixture below.
        with tempfile.TemporaryDirectory() as tmp:
            path=Path(tmp)
            write_json(path/'state.json', {'id':'b'*32,'started_at':now(),'status':'queued','config':{},'metadata':{},'files':[], 'steps':{},'events':[],'timings':[],'feedback':'','model_calls':[],'answers':[],'questions':[]})
            FakeProvider.calls=[]
            run=Runner(path,provider_factory=FakeProvider,exporter=lambda *args: {})
            run.execute()
            self.assertEqual(run.state['status'],'completed')
            qid=Runner.safe_id(PROBLEM['questions'][0])
            self.assertLess(FakeProvider.calls.index(qid+'-review-plan-1'),FakeProvider.calls.index(qid+'-revision-1'))
            self.assertLess(FakeProvider.calls.index(qid+'-revision-1'),FakeProvider.calls.index(qid+'-review-plan-2'))
            self.assertIn('适用条件',run.state['steps']['revision2-'+qid]['review_plan']['directions'][0]['reason'])
