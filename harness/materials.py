from __future__ import annotations

from collections import Counter
import csv
import io
import re
from pathlib import Path
from zipfile import ZipFile
from xml.etree import ElementTree as ET

NS = {'w': 'http://schemas.openxmlformats.org/wordprocessingml/2006/main'}
IMAGES = {'.png', '.jpg', '.jpeg', '.webp'}
SUPPORTED = IMAGES | {'.docx', '.pdf', '.txt', '.md', '.csv', '.tsv', '.xlsx'}


def chunk_text(text, name, prefix, locator):
    parts, current = [], ''
    for line in text.splitlines():
        for start in range(0, max(1, len(line)), 2200):
            piece = line[start:start + 2200]
            if len(current) + len(piece) > 2500:
                parts.append(current)
                current = ''
            current += piece + '\n'
    if current.strip():
        parts.append(current)
    return [{'id': f'{prefix}-{i+1}', 'name': name, 'locator': locator + f' 片段 {i+1}', 'text': p}
            for i, p in enumerate(parts)]


def extract(path, source_id, working):
    path, working = Path(path), Path(working)
    working.mkdir(parents=True, exist_ok=True)
    suffix = path.suffix.lower()
    if suffix not in SUPPORTED:
        raise ValueError(f'不支持 {suffix}；请转换为 PDF、DOCX、XLSX、文本或图片。')
    chunks, visuals = [], []
    if suffix in IMAGES:
        visuals.append((path, source_id, '图片'))
    elif suffix == '.docx':
        with ZipFile(path) as z:
            root = ET.fromstring(z.read('word/document.xml'))
            lines = []
            for node in root.find('w:body', NS):
                if node.tag.endswith('}tbl'):
                    for row in node.findall('w:tr', NS):
                        lines.append(' | '.join(' '.join(t.text or '' for t in cell.findall('.//w:t', NS))
                                                for cell in row.findall('w:tc', NS)))
                elif node.tag.endswith('}p'):
                    text = ''.join(t.text or '' for t in node.findall('.//w:t', NS))
                    if text:
                        lines.append(text)
            chunks += chunk_text('\n'.join(lines), path.name, source_id, '正文')
            for i, name in enumerate(z.namelist()):
                if name.startswith('word/media/') and Path(name).suffix.lower() in IMAGES:
                    dest = working / (source_id + '-' + str(i) + Path(name).suffix)
                    dest.write_bytes(z.read(name))
                    visuals.append((dest, source_id + 'v' + str(i), '内嵌图片 ' + Path(name).name))
            vector_media = [name for name in z.namelist() if name.startswith('word/media/') and
                            Path(name).suffix.lower() in ('.emf', '.wmf', '.svg')]
        if vector_media:
            from .common import runtime, skill_dir
            from .exports import command, font_env
            preview = working / (source_id + '-document-preview')
            command([runtime('python'), skill_dir('documents') / 'render_docx.py', path,
                     '--output_dir', preview], working, timeout=360, env=font_env(working))
            for page in sorted(preview.glob('page-*.png'), key=lambda p: int(p.stem.split('-')[-1])):
                visuals.append((page, source_id + 'render' + page.stem.split('-')[-1], '文档渲染 ' + page.stem))
    elif suffix == '.pdf':
        from pypdf import PdfReader
        import pypdfium2 as pdfium
        reader, render = PdfReader(str(path)), None
        for i, page in enumerate(reader.pages):
            text = page.extract_text() or ''
            chunks += chunk_text(text, path.name, source_id + 'p' + str(i+1), f'第 {i+1} 页')
            content = page.get_contents()
            operations = content.get_data() if content else b''
            vector_drawing = len(re.findall(rb'\s(?:m|l|c|re)\s', operations)) > 12
            if len(text.strip()) < 60 or len(page.images) or vector_drawing:
                if render is None:
                    render = pdfium.PdfDocument(str(path))
                dest = working / f'{source_id}-page-{i+1}.png'
                render[i].render(scale=1.5).to_pil().save(dest)
                visuals.append((dest, source_id + 'p' + str(i+1) + 'v', f'第 {i+1} 页图像'))
        if render:
            render.close()
    elif suffix == '.xlsx':
        import openpyxl
        book = openpyxl.load_workbook(path, data_only=True, read_only=True)
        formulas = openpyxl.load_workbook(path, data_only=False, read_only=True)
        try:
            for i, sheet in enumerate(book):
                rows = []
                for values, originals in zip(sheet.iter_rows(values_only=True), formulas.worksheets[i].iter_rows(values_only=True)):
                    if any(isinstance(o,str) and o.startswith('=') and v is None for v,o in zip(values, originals)):
                        raise ValueError(path.name + ' 中的公式没有保存计算结果，请在 Excel 中重新计算后上传，或提供计算后的 CSV。')
                    rows.append(' | '.join('' if c is None else str(c) for c in values))
                chunks += chunk_text('\n'.join(rows), path.name, source_id + 's' + str(i+1), sheet.title)
        finally:
            book.close()
            formulas.close()
    else:
        raw = path.read_bytes()
        text = None
        encodings = ('utf-16',) if raw.startswith((b'\xff\xfe', b'\xfe\xff')) else ('utf-8-sig', 'gb18030')
        for encoding in encodings:
            try:
                text = raw.decode(encoding)
                break
            except UnicodeError:
                pass
        if text is None:
            raise ValueError(f'{path.name} 无法按文本读取，请转换为 UTF-8。')
        chunks += chunk_text(text, path.name, source_id, '正文')
    if not chunks and not visuals:
        raise ValueError(path.name + ' 没有可读取内容。')
    return chunks, visuals


def terms(text):
    english = re.findall(r'[a-z0-9]{2,}', text.lower())
    chinese = re.findall(r'[\u4e00-\u9fff]+', text)
    return Counter(english + [s[i:i+2] for s in chinese for i in range(len(s)-1)])


def select(chunks, query, limit=65000):
    """Keep full local source chunks; select relevant passages without rewriting them."""
    q = terms(query)
    ranked = sorted(enumerate(chunks), key=lambda item: sum((terms(item[1]['text']) & q).values()), reverse=True)
    picked, used = {}, 0
    for index, chunk in ranked:
        length = len(chunk['text'])
        if used + length <= limit:
            picked[index] = chunk
            used += length
    return [picked[k] for k in sorted(picked)]
