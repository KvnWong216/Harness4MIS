"""Shared MIS303 Word styles, native prompt formatting and Part selection."""
import re
from docx.shared import Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

FONTS = {0: '方正小标宋简体', 1: '黑体', 2: '楷体', 3: '仿宋', 4: '仿宋', 5: '仿宋'}
CN = '一二三四五六七八九十'

def chinese(n):
    if n <= 10: return CN[n-1]
    return (CN[n//10-1] if n//10 > 1 else '') + '十' + (CN[n%10-1] if n%10 else '')

def font(target, east='仿宋', size=16):
    target.font.name = 'Times New Roman'
    target.font.size = Pt(size)
    target.font.color.rgb = RGBColor(0, 0, 0)
    el = target.element if hasattr(target, 'element') else target._element
    rf = el.get_or_add_rPr().get_or_add_rFonts()
    for k in ['ascii', 'hAnsi', 'cs']: rf.set(qn('w:'+k), 'Times New Roman')
    rf.set(qn('w:eastAsia'), east)
    for k in ['asciiTheme', 'hAnsiTheme', 'eastAsiaTheme', 'cstheme']:
        rf.attrib.pop(qn('w:'+k), None)

def paragraph(p, level=None, title=False, table=False):
    f = p.paragraph_format
    f.line_spacing = Pt(35 if title else 29)
    f.space_before = Pt(0); f.space_after = Pt(0)
    f.first_line_indent = Pt(0 if title or level or table else 32)
    ppr = p._p.get_or_add_pPr()
    ind = ppr.get_or_add_ind()
    if title or level or table: ind.attrib.pop(qn('w:firstLineChars'), None)
    else: ind.set(qn('w:firstLineChars'), '200')
    grid = ppr.find(qn('w:snapToGrid'))
    if grid is None: grid = OxmlElement('w:snapToGrid'); ppr.append(grid)
    grid.set(qn('w:val'), '0')
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER if title else (WD_ALIGN_PARAGRAPH.LEFT if level else WD_ALIGN_PARAGRAPH.JUSTIFY)
    if title or level: f.keep_with_next = True
    for run in p.runs: font(run, FONTS[0 if title else level] if title or level else '仿宋', 22 if title else 16)
    return p

def configure(doc):
    for name, level in [('Normal', None), ('Title', 0), *[('Heading '+str(i), i) for i in range(1,6)]]:
        style = doc.styles[name]
        font(style, FONTS[level] if level is not None else '仿宋', 22 if level == 0 else 16)
        style.font.bold = level in (1,2)
        f=style.paragraph_format; f.line_spacing=Pt(35 if level == 0 else 29)
        f.space_after=Pt(0); f.space_before=Pt(0); f.first_line_indent=Pt(32 if level is None else 0)
        f.keep_with_next=level is not None
    for root in [doc.styles.element, doc._element]:
        for border in root.xpath('.//w:pBdr'): border.getparent().remove(border)
    COUNTERS[id(doc)]=[0]*5
    return doc

def heading(doc, text, level=1):
    if not 1 <= level <= 5: raise ValueError('Heading level must be 1–5')
    # python-docx Document uses slots, so counters are stored in a module map.
    counts=COUNTERS.setdefault(id(doc), [0]*5)
    counts[level-1]+=1
    for i in range(level,5): counts[i]=0
    n=counts[level-1]
    text=re.sub(r'^\s*(?:[一二三四五六七八九十]+、|（[一二三四五六七八九十]+）|\d+[.、]\s*|（\d+）|[①②③④⑤⑥⑦⑧⑨⑩])\s*','',text)
    prefix={1:chinese(n)+'、',2:'（'+chinese(n)+'）',3:str(n)+'.',4:'（'+str(n)+'）',5:chr(0x2460+n-1)}[level]
    return paragraph(doc.add_paragraph(prefix+text, 'Heading '+str(level)), level=level)

COUNTERS={}

def inline(p, text):
    # A small, deterministic Markdown subset: emphasis, code text and links.
    pattern=r'(\*\*.*?\*\*|__.*?__|\*[^*]+\*|_[^_]+_|`[^`]+`|\[[^]]+\]\([^)]+\))'
    for token in re.split(pattern,text):
        if not token: continue
        bold=token.startswith(('**','__')); italic=not bold and token.startswith(('*','_'))
        if bold: value=token[2:-2]
        elif italic or token.startswith('`'): value=token[1:-1]
        elif re.match(r'\[[^]]+\]\([^)]+\)', token):
            m=re.match(r'\[([^]]+)\]\(([^)]+)\)',token); value=m[1]+' ('+m[2]+')'
        else: value=token
        r=p.add_run(value);font(r);r.bold=bold;r.italic=italic
    return p

def markdown(doc, text, base_level=2):
    fenced=False
    for line in text.splitlines():
        line=line.strip()
        if line.startswith('```'): fenced=not fenced;continue
        if not line: continue
        match=re.match(r'^#{1,6}\s*(.+)$',line) if not fenced else None
        if match:
            heading(doc,match[1],min(base_level,5))
        else:
            p=doc.add_paragraph();inline(p,line);paragraph(p)

PART_RE=re.compile(r'(?im)^.*?(?:Assign[^\n]*?[- ]|^)(Part\s*[- ]?\d+)[^\n]*$')
def part_key(value):
    m=re.search(r'(?:part\s*[- ]?)?(\d+)',str(value),re.I)
    return m[1] if m else ''

def select_part(text, part):
    matches=list(PART_RE.finditer(text)); wanted=part_key(part)
    if not wanted or not matches: return text
    blocks=[]
    for i,m in enumerate(matches):
        if part_key(m[1])==wanted:
            blocks.append(text[m.start():matches[i+1].start() if i+1<len(matches) else len(text)].strip())
    if not blocks: raise ValueError('题目中没有找到所选 Part '+wanted)
    return '\n\n'.join(blocks)

def teacher_prompts(text):
    blocks=re.findall(r'【\s*(#Role[\s\S]*?)】',text,re.I)
    if blocks: return blocks
    start=re.search(r'(?im)^#\s*Role\s*[:：]',text)
    return [text[start.start():].strip()] if start else []

def requirements(doc,text):
    for line in text.splitlines():
        line=line.strip()
        if not line or line in ['【','】']: continue
        if re.match(r'(?i)^#\s*(Role|Background|Task|Requirement|附件)',line):
            heading(doc,re.sub(r'^#\s*','',line),2)
        elif re.match(r'(?i)^Assign.*Part|^.*阶段工作任务',line): heading(doc,line,2)
        else:
            p=doc.add_paragraph();inline(p,line);paragraph(p)
